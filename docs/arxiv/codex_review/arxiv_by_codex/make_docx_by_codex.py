"""Build the editable Word version of the arXiv manuscript from main_by_codex.tex.

    /home/mirlab/anaconda3/envs/flow_matching/bin/python make_docx.py

Two stages:
  1. Preprocess main_by_codex.tex -> _by_codex_docx.tex: resolve \\cite / \\ref to literal numbers (pandoc has
     no .aux), switch figures to PNG, turn the title block and thebibliography into plain
     paragraphs. Section numbering is emitted literally (I, II.A, ...) so the docx matches
     the PDF.
  2. pandoc 3 (texbuild env) -> main_by_codex.docx, then python-docx post-pass: 1-inch margins,
     Times New Roman, black headings, and CENTERED table cells (all data columns; first
     column stays left) with a fixed layout so columns cannot collapse.

The \\ref map below is maintained BY HAND -- if a section/table/figure is added or reordered
in main_by_codex.tex, update it (the script asserts every \\ref it sees is in the map, so a stale map
fails loudly rather than printing '??').
"""
import re
import subprocess
import sys

PANDOC = "/home/mirlab/anaconda3/envs/texbuild/bin/pandoc"

REF = {
    "sec:intro": "I", "sec:method": "II", "sec:forward": "II.A", "sec:bridge": "II.B",
    "sec:loop": "II.C", "sec:impl": "II.D", "sec:setup": "III", "sec:data": "III.A",
    "sec:protocol": "III.B", "sec:baselines": "III.C", "sec:results": "IV",
    "sec:results-main": "IV.A", "sec:results-cmp": "IV.B", "sec:cost": "IV.C",
    "sec:discussion": "V", "sec:conclusion": "VI",
    "tab:main": "1", "tab:cmp": "2", "tab:dof": "3", "tab:cost": "4",
    "fig:manifold": "1", "fig:method": "2", "fig:images": "3", "fig:quant": "4",
    "fig:bench": "5", "fig:motion": "6",
    "eq:objective": "1", "eq:bridge": "2", "eq:tangent": "3",
    "alg:loop": "1",
}
SECNUM = {  # literal numbers prepended to headings, in document order
    "Introduction": "I.", "Methods": "II.", "Experiments": "III.", "Results": "IV.",
    "Discussion": "V.", "Conclusion": "VI.",
    "Problem formulation": "A.", "Flow-matching prior on the geometry bridge": "B.",
    "Predictor-corrector inference loop": "C.", "Implementation details": "D.",
    "Data and simulation setup": "A.", "Evaluation protocol": "B.",
    "Comparison methods": "C.",
    "Performance of the proposed method": "A.", "Comparison with other methods": "B.", "Computational cost": "C.",
}


def main():
    src = open("main_by_codex.tex").read()

    # ---- cite map from thebibliography order
    keys = re.findall(r"\\bibitem\{(\w+)\}", src)
    cite_no = {k: i + 1 for i, k in enumerate(keys)}

    def cite(m):
        nums = [str(cite_no[k.strip()]) for k in m.group(1).split(",")]
        return "[" + ",".join(nums) + "]"

    body = src
    body = re.sub(r"~?\\cite\{([^}]+)\}", cite, body)

    def ref(m):
        k = m.group(1)
        assert k in REF, f"\\ref{{{k}}} missing from REF map -- update make_docx.py"
        return REF[k]

    body = re.sub(r"\\ref\{([^}]+)\}", ref, body)
    body = re.sub(r"\\eqref\{([^}]+)\}", lambda m: "(" + ref(m) + ")", body)
    body = body.replace(".pdf}", ".png}")
    # `$^\circ$` (empty-base superscript) renders as a placeholder box in Word; use the
    # plain degree character in text positions.
    body = body.replace("$^\\circ$", "\u00b0")

    # Number the captions ("Table N: ...", "Figure N: ...") in document order -- pandoc's
    # docx writer emits the caption text alone, while the body cites them by number.
    counters = {"table": 0, "figure": 0}
    pieces, pos = [], 0
    for m in re.finditer(r"\\begin\{(table|figure)\}.*?\\end\{\1\}", body, re.S):
        env = m.group(1)
        counters[env] += 1
        label = ("Table" if env == "table" else "Figure") + f" {counters[env]}: "
        chunk = m.group(0).replace("\\caption{", "\\caption{" + label, 1)
        pieces.append(body[pos:m.start()])
        pieces.append(chunk)
        pos = m.end()
    pieces.append(body[pos:])
    body = "".join(pieces)

    # ---- algorithm floats -> a bold "Algorithm N" line + a one-column table (pandoc has no
    # algorithmic support; a table keeps the math and the line structure in Word)
    def algo(m):
        block = m.group(0)
        cap = re.search(r"\\caption\{(.*?)\}", block).group(1)
        body_ = re.search(r"\\begin\{algorithmic\}(?:\[\d\])?(.*?)\\end\{algorithmic\}",
                          block, re.S).group(1)
        rows, depth, n = [], 0, 0
        for ln in body_.strip().splitlines():
            ln = ln.strip()
            if not ln:
                continue
            ln = re.sub(r"\\Comment\{(.*?)\}$", r"  \\textit{(\1)}", ln)
            ln = ln.replace(",\\quad", ",\u2002").replace("\\quad", "\u2002")
            if ln.startswith("\\Require"):
                txt = "\\textbf{Input:} " + ln[len("\\Require"):].strip()
            elif ln.startswith("\\EndFor"):
                depth -= 1; txt = "\\textbf{end for}"
            elif ln.startswith("\\For"):
                cond = re.match(r"\\For\{(.*)\}$", ln).group(1)
                txt = "\\textbf{for} " + cond + " \\textbf{do}"
            elif ln.startswith("\\State \\Return"):
                txt = "\\textbf{return} " + ln[len("\\State \\Return"):].strip()
            elif ln.startswith("\\State"):
                txt = ln[len("\\State"):].strip()
            else:
                txt = ln
            if ln.startswith("\\Require"):
                rows.append(txt + " \\\\")
                continue
            n += 1
            rows.append(f"{n}: " + "\u2003" * depth + txt + " \\\\")
            if ln.startswith("\\For"):
                depth += 1
        head = f"\\textbf{{Algorithm 1: {cap}}}\n\n"
        return head + "\\begin{tabular}{l}\n\\hline\n" + "\n".join(rows) + "\n\\hline\n\\end{tabular}\n"

    body = re.sub(r"\\begin\{algorithm\}.*?\\end\{algorithm\}", algo, body, flags=re.S)

    # ---- title block -> plain paragraphs (pandoc drops \maketitle content otherwise)
    title = re.search(r"\\title\{\\bfseries\s*(.*?)\}\n", body, re.S).group(1)
    title = " ".join(title.split())
    body = re.sub(r"\\title\{.*?\n\n", "", body, flags=re.S)
    body = re.sub(r"\\author\{.*?\}\}\n", "", body, flags=re.S)
    body = body.replace("\\date{}", "")
    author_block = (
        "\\begin{center}{\\Large \\textbf{" + title + "}}\\\\[10pt]\n"
        "Sungho Yun and Seungryong Cho\\\\\n"
        "Department of Nuclear and Quantum Engineering,\\\\\n"
        "Korea Advanced Institute of Science and Technology (KAIST), "
        "Daejeon 34141, South Korea\\end{center}\n\n"
        "\\textbf{Abstract}\n"
    )
    body = body.replace("\\maketitle\n", author_block)
    body = body.replace("\\begin{abstract}\n\\noindent\n", "")
    body = body.replace("\\end{abstract}\n", "")

    # ---- literal section numbers (pandoc numbering would restart per style)
    def sec(m):
        name = m.group(2)
        num = SECNUM.get(name, "")
        return f"\\{m.group(1)}{{{num} {name}}}" if num else m.group(0)

    body = re.sub(r"\\(section|subsection)\*?\{([^}]*)\}", sec, body)
    body = re.sub(r"\\label\{[^}]*\}", "", body)

    # ---- bibliography -> plain numbered paragraphs
    bib = re.search(r"\\begin\{thebibliography\}.*\\end\{thebibliography\}", body, re.S)
    items = re.split(r"\\bibitem\{\w+\}", bib.group(0))[1:]
    out = ["\\section*{References}"]
    for i, it in enumerate(items, 1):
        it = it.replace("\\end{thebibliography}", "").strip()
        it = " ".join(it.split())
        out.append(f"[{i}] {it}\n")
    body = body.replace(bib.group(0), "\n\n".join(out))

    open("_by_codex_docx.tex", "w").write(body)
    subprocess.run([PANDOC, "_by_codex_docx.tex", "-o", "main_by_codex_raw.docx",
                    "--resource-path=."], check=True)

    # ---- python-docx post-pass ---------------------------------------------------------
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document("main_by_codex_raw.docx")
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(1)
        s.left_margin = s.right_margin = Inches(1)

    for st in doc.styles:
        if st.type is not None and getattr(st, "font", None) is not None:
            try:
                st.font.name = "Times New Roman"
            except Exception:
                pass
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    for st in doc.styles:
        if st.name and st.name.startswith("Heading"):
            st.font.color.rgb = RGBColor(0, 0, 0)
            st.font.name = "Times New Roman"

    # Tables: real grid (the pandoc->Word collapse trap), fixed layout, centred data cells.
    page_in = 6.5
    for tbl in doc.tables:
        ncol = len(tbl.rows[0].cells)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblPr = tbl._tbl.tblPr
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tblPr.append(layout)
        for g in tbl._tbl.findall(qn("w:tblGrid")):
            tbl._tbl.remove(g)
        grid = OxmlElement("w:tblGrid")
        first = page_in if ncol == 1 else page_in * 0.34
        rest = (page_in - first) / max(ncol - 1, 1)
        for j in range(ncol):
            gc = OxmlElement("w:gridCol")
            w = first if j == 0 else rest
            gc.set(qn("w:w"), str(int(w * 1440)))
            grid.append(gc)
        tbl._tbl.insert(list(tbl._tbl).index(tblPr) + 1, grid)
        for r in tbl.rows:
            for j, c in enumerate(r.cells):
                for p in c.paragraphs:
                    p.alignment = (WD_ALIGN_PARAGRAPH.LEFT if j == 0
                                   else WD_ALIGN_PARAGRAPH.CENTER)
                    for run in p.runs:
                        run.font.size = Pt(10)
                        run.font.name = "Times New Roman"

    # Centre images and their captions.
    for p in doc.paragraphs:
        if p.style.name in ("Captioned Figure", "Image Caption", "Caption") or \
           any(r._element.findall(qn("w:drawing")) for r in p.runs):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Headings: run-level Times overrides the theme's sans headings font; black, bold.
    hsize = {"Heading 1": 14, "Heading 2": 12.5, "Heading 3": 12}
    for p in doc.paragraphs:
        if p.style.name in hsize:
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(hsize[p.style.name])
                run.font.bold = True
                run.font.color.rgb = RGBColor(0, 0, 0)
                rpr = run._element.get_or_add_rPr()
                rf = rpr.find(qn("w:rFonts"))
                if rf is None:
                    rf = OxmlElement("w:rFonts")
                    rpr.append(rf)
                for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                    rf.set(qn(attr), "Times New Roman")

    doc.save("main_by_codex.docx")
    print("wrote main_by_codex.docx")


if __name__ == "__main__":
    sys.exit(main())
